import sqlite3
from pathlib import Path

import openpyxl


root = Path(__file__).resolve().parents[1]
workbook_path = root / "data" / "verification" / "20260830-历史数据整理清单.xlsx"

with sqlite3.connect(root / "data" / "app.db") as db:
    industries = dict(
        db.execute(
            "SELECT id, industry FROM customers WHERE id LIKE 'legacy-customer-%' AND industry IS NOT NULL AND trim(industry) != ''"
        )
    )

workbook = openpyxl.load_workbook(workbook_path)
sheet = workbook["客户行业与代理商确认"]
updated = 0
for row in range(5, sheet.max_row + 1):
    industry = industries.get(sheet.cell(row, 1).value)
    if industry and sheet.cell(row, 4).value == "缺失":
        sheet.cell(row, 3).value = industry
        sheet.cell(row, 4).value = "已自动补齐"
        sheet.cell(row, 8).value = industry
        updated += 1

replacements = {
    "已自动补齐54家，当前缺54家": "已自动补齐85家，当前缺23家",
    "累计已自动补齐54家，当前仍缺54家": "累计已自动补齐85家，当前仍缺23家",
}
for current_sheet in workbook.worksheets:
    for row in current_sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                for old, new in replacements.items():
                    cell.value = cell.value.replace(old, new)

workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.calculation.calcMode = "auto"
workbook.save(workbook_path)
assert updated == 31, f"预期同步31家，实际{updated}家"
print({"updated": updated, "workbook": str(workbook_path)})
