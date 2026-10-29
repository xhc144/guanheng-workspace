"""风险、问题、看板、导入导出"""
from decimal import Decimal
from datetime import date, datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy.orm import Session
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from pydantic import ValidationError
from typing import Optional
from io import BytesIO
from app.database import get_db
from app.models import (
    Risk, RiskFollowUp, Issue, Project, ProjectMember, Task, Receipt, Customer,
    Contract, ActivityLog, AuditLog, Notification, User
)
from app.schemas import RiskCreate, RiskFollowUpCreate, DashboardData, ReceiptCreate
from app.auth import get_current_user, require_role
from app.routes.finance import _add_receipt, _lock_receipt_write
from app.routes.project_access import is_project_admin, require_project, scope_project_rows, visible_project_ids

risk_router = APIRouter(prefix="/api/risks", tags=["风险"])


@risk_router.get("")
def list_risks(project_id: Optional[str] = Query(None), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    q = scope_project_rows(db.query(Risk), Risk.project_id, current_user)
    if project_id:
        require_project(db, current_user, project_id)
        q = q.filter(Risk.project_id == project_id)
    return q.order_by(Risk.created_at.desc()).all()


@risk_router.post("")
def create_risk(req: RiskCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    require_project(db, current_user, req.project_id)
    # 自动计算风险等级
    prob_map = {"low": 1, "medium": 2, "high": 3}
    impact_map = {"low": 1, "medium": 2, "high": 3, "critical": 4}
    score = prob_map.get(req.probability, 2) * impact_map.get(req.impact, 2)
    if score <= 3:
        level = "low"
    elif score <= 6:
        level = "medium"
    elif score <= 9:
        level = "high"
    else:
        level = "critical"
    level_reason = f"概率={req.probability}({prob_map.get(req.probability,2)}) x 影响={req.impact}({impact_map.get(req.impact,2)}) = 评分{score} → 等级{level}"
    r = Risk(**req.model_dump(), risk_level=level, level_reason=level_reason)
    db.add(r)
    db.add(ActivityLog(project_id=r.project_id, user_id=current_user.id, action="登记风险", detail=r.title))
    db.flush()
    # 通知负责人
    if r.responsible_id:
        db.add(Notification(user_id=r.responsible_id, title=f"新风险分配: {r.title}", content=f"项目中的风险已分配给您，等级: {level}", ntype="risk", link=f"/projects/{r.project_id}"))
    db.commit()
    db.refresh(r)
    return r


def _parse_commitment_date(value):
    """接受 datetime 或 ISO 日期字符串，返回 datetime。"""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
    return None


@risk_router.post("/{risk_id}/follow-up")
def add_follow_up(risk_id: str, req: RiskFollowUpCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    risk = db.query(Risk).filter(Risk.id == risk_id).first()
    if not risk:
        raise HTTPException(status_code=404, detail="风险不存在")
    require_project(db, current_user, risk.project_id)
    fu = RiskFollowUp(
        risk_id=risk_id,
        content=req.content,
        commitment_date=_parse_commitment_date(req.commitment_date),
        created_by=current_user.id,
    )
    db.add(fu)
    db.add(ActivityLog(project_id=risk.project_id, user_id=current_user.id, action="风险跟进", detail=req.content[:100]))
    db.commit()
    db.refresh(fu)
    return fu


@risk_router.delete("/{risk_id}")
def delete_risk(risk_id: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    r = db.query(Risk).filter(Risk.id == risk_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="风险不存在")
    require_project(db, current_user, r.project_id)
    db.delete(r)
    db.commit()
    return {"ok": True}


dashboard_router = APIRouter(prefix="/api/dashboard", tags=["看板"])


@dashboard_router.get("")
def get_dashboard(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project_ids = visible_project_ids(current_user)
    now = datetime.utcnow()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    total_projects = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.is_archived == False).count()
    active_projects = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.status == "active", Project.is_archived == False).count()
    total_contract_amount = db.query(func.coalesce(func.sum(Contract.amount), 0)).filter(Contract.project_id.in_(project_ids)).scalar() or Decimal("0")
    total_received = db.query(func.coalesce(func.sum(Receipt.amount), 0)).filter(Receipt.project_id.in_(project_ids)).filter(Receipt.status == "confirmed").scalar() or Decimal("0")
    total_invoiced = db.query(func.coalesce(func.sum(Contract.amount), 0)).filter(Contract.project_id.in_(project_ids)).scalar() or Decimal("0")
    total_receivable = total_invoiced - total_received
    overdue_projects = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.planned_end < now, Project.status == "active").count()
    high_risk_projects = scope_project_rows(db.query(Risk), Risk.project_id, current_user).filter(Risk.risk_level.in_(["high", "critical"]), Risk.status == "open").count()

    # 月度回款趋势（最近12个月）
    monthly_receipts = []
    for i in range(11, -1, -1):
        m_start = datetime(now.year, now.month, 1) - timedelta(days=i * 30)
        m_start = m_start.replace(day=1)
        if i == 0:
            m_end = now
        else:
            next_m = m_start.replace(day=28) + timedelta(days=4)
            m_end = next_m.replace(day=1) - timedelta(days=1)
        amt = db.query(func.coalesce(func.sum(Receipt.amount), 0)).filter(Receipt.project_id.in_(project_ids)).filter(
            Receipt.received_date >= m_start, Receipt.received_date <= m_end, Receipt.status == "confirmed"
        ).scalar() or Decimal("0")
        monthly_receipts.append({"month": m_start.strftime("%Y-%m"), "amount": float(amt)})

    # 项目健康度分布
    health_dist = []
    for h in ["green", "yellow", "red"]:
        cnt = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.health == h, Project.is_archived == False).count()
        health_dist.append({"health": h, "count": cnt})

    # TOP风险
    from sqlalchemy import case
    top_risks = []
    for r in scope_project_rows(db.query(Risk), Risk.project_id, current_user).filter(Risk.status == "open").order_by(
        case({"critical": 0, "high": 1, "medium": 2, "low": 3}, value=Risk.risk_level)
    ).limit(10).all():
        top_risks.append({"id": r.id, "title": r.title, "level": r.risk_level, "project_id": r.project_id})

    # 应收账龄（基于未收款发票/未核销合同）
    aging_0_30 = aging_31_60 = aging_61_90 = aging_90_plus = 0.0
    from app.models import Invoice
    unpaid = db.query(Invoice).filter(Invoice.contract.has(Contract.project_id.in_(project_ids))).filter(Invoice.status.in_(["issued", "partial_paid"])).all()
    for inv in unpaid:
        days = (now - (inv.due_date or inv.issued_date or now)).days
        amt = float(inv.amount or 0)
        if days <= 30: aging_0_30 += amt
        elif days <= 60: aging_31_60 += amt
        elif days <= 90: aging_61_90 += amt
        else: aging_90_plus += amt

    # 项目阶段分布
    stage_dist = []
    for stage in db.query(Project.stage, func.count(Project.id)).filter(Project.id.in_(project_ids)).filter(Project.is_archived == False).group_by(Project.stage).all():
        stage_dist.append({"stage": stage[0] or "未设置", "count": stage[1]})

    # 项目状态分布
    status_dist = []
    for status, cnt in db.query(Project.status, func.count(Project.id)).filter(Project.id.in_(project_ids)).filter(Project.is_archived == False).group_by(Project.status).all():
        status_dist.append({"status": status or "未设置", "count": cnt})

    return {
        "total_projects": total_projects,
        "active_projects": active_projects,
        "total_contract_amount": float(total_contract_amount),
        "total_received": float(total_received),
        "total_receivable": float(total_receivable),
        "total_budget": 0,
        "total_cost": 0,
        "overdue_projects": overdue_projects,
        "high_risk_projects": high_risk_projects,
        "monthly_receipts": monthly_receipts,
        "project_health_dist": health_dist,
        "top_risks": top_risks,
        "aging_0_30": aging_0_30,
        "aging_31_60": aging_31_60,
        "aging_61_90": aging_61_90,
        "aging_90_plus": aging_90_plus,
        "stage_dist": stage_dist,
        "status_dist": status_dist,
        "top_overdue": [],
        "top_receivable": [],
    }


@dashboard_router.get("/summary")
def get_dashboard_summary(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project_ids = visible_project_ids(current_user)
    """仪表盘摘要：仅返回关键KPI"""
    from decimal import Decimal
    now = datetime.utcnow()
    active = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.status == "active", Project.is_archived == False).count()
    receipts = db.query(func.coalesce(func.sum(Receipt.amount), 0)).filter(Receipt.project_id.in_(project_ids)).filter(Receipt.status == "confirmed").scalar() or Decimal("0")
    risks = scope_project_rows(db.query(Risk), Risk.project_id, current_user).filter(Risk.status == "open", Risk.risk_level.in_(["high", "critical"])).count()
    overdue = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.planned_end < now, Project.status == "active").count()
    return {
        "active_projects": active,
        "total_received": float(receipts),
        "high_risk_count": risks,
        "overdue_projects": overdue,
        "generated_at": now.isoformat(),
    }


import_router = APIRouter(prefix="/api/import", tags=["导入"])


def _excel_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _excel_amount(value: object) -> str:
    # Keep the existing display-format support; ReceiptCreate validates the
    # resulting value without rounding, clamping or substituting zero.
    return _excel_text(value).replace(",", "").replace("¥", "").replace(" ", "")


def _excel_received_date(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    try:
        return datetime.fromisoformat(_excel_text(value))
    except ValueError:
        raise ValueError("回款日期格式错误，请使用 YYYY-MM-DD 或 Excel 日期单元格")


def _excel_error(exc: Exception, amount_label: str) -> str:
    if isinstance(exc, ValidationError):
        messages = {
            "greater_than": "必须大于零",
            "finite_number": "必须为有限数值",
            "decimal_parsing": "格式错误",
            "decimal_max_places": "最多保留两位小数",
            "decimal_max_digits": "总位数不得超过18位",
            "decimal_whole_digits": "整数部分不得超过16位",
        }
        return "；".join(
            f"{amount_label if error['loc'] == ('amount',) else '回款字段'}: "
            f"{messages.get(error['type'], error['msg'])}"
            for error in exc.errors()
        )
    if isinstance(exc, HTTPException):
        return str(exc.detail)
    if isinstance(exc, IntegrityError):
        return "数据约束冲突，请核对该行"
    return str(exc)


def _initial_receipt_request(row: dict, project_id: str, customer_name: str) -> Optional[ReceiptCreate]:
    raw = _excel_amount(row.get("初始已收金额"))
    if not raw:
        return None
    try:
        amount = Decimal(raw)
    except ArithmeticError:
        raise ValueError("初始已收金额格式错误")
    # Blank and explicit zero mean no opening receipt in a project import.
    # Every other value must pass the same positive-money schema as JSON.
    if amount.is_finite() and amount.is_zero():
        return None
    return ReceiptCreate(
        project_id=project_id, amount=amount, received_date=datetime.utcnow(),
        payer=customer_name, bank_ref=f"IMPORT-{_excel_text(row.get('项目编号'))}",
        notes="Excel 初始已收金额导入",
    )


def _save_excel_receipt(req: ReceiptCreate, db: Session, current_user: User, *, confirmed: bool = False) -> Receipt:
    receipt = _add_receipt(req, db, current_user)
    if confirmed:
        receipt.status = "confirmed"
        receipt.confirmed_by = current_user.id
        receipt.confirmed_at = datetime.utcnow()
        db.add(AuditLog(user_id=current_user.id, action="confirm", target_type="receipt",
                        target_id=receipt.id, new_value=str(req.amount)))
    db.flush()
    db.refresh(receipt, attribute_names=["amount"])
    # SQLite Numeric uses a floating-point binding. Even a schema-valid amount
    # near its maximum can lose cents; reject that row instead of saving a
    # different amount. The caller's savepoint includes all logs and related rows.
    if receipt.amount != req.amount:
        label = "初始已收金额" if confirmed else "回款金额"
        raise ValueError(f"{label}无法按两位小数精确保存，请核对金额")
    return receipt


def _project_import_rows(wb) -> Optional[list[tuple[int, dict]]]:
    rows, headers, header_row, best_size = [], [], 0, 0
    for ws in wb.worksheets:
        sheet_rows = list(ws.iter_rows(values_only=True))
        for pos, candidate in enumerate(sheet_rows[:10]):
            candidate_headers = [_excel_text(h) for h in candidate]
            if ("项目编号" in candidate_headers and "客户名称" in candidate_headers
                    and len(sheet_rows) > best_size):
                rows, headers, header_row, best_size = sheet_rows, candidate_headers, pos, len(sheet_rows)
                break
    if not headers:
        return None
    detail = []
    for idx, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        values = dict(zip(headers, row))
        code = _excel_text(values.get("项目编号"))
        name = _excel_text(values.get("项目名称")) or _excel_text(values.get("客户名称"))
        if not any(_excel_text(value) for value in row):
            continue
        if code.lower() in ("合计", "总计", "total") or "合计" in name or "总计" in name:
            continue
        detail.append((idx, values))
    return detail


def _import_project_rows(wb, db: Session, current_user: User, errors: list[str]) -> int:
    rows = _project_import_rows(wb)
    if rows is None:
        errors.append("未找到项目明细表头：需要项目编号和客户名称列")
        return 0

    # Preserve the existing whole-batch authorization preflight. Never turn a
    # forbidden target or confirmation into a partial import.
    for _, row in rows:
        code = _excel_text(row.get("项目编号"))
        target = db.query(Project).filter(Project.code == code).first()
        if target:
            require_project(db, current_user, target.id)
            if target.manager_id != current_user.id and not is_project_admin(current_user):
                raise HTTPException(status_code=403, detail="只有项目负责人或管理员可导入项目归属修改")
        try:
            opening = _initial_receipt_request(row, "pending", _excel_text(row.get("客户名称")))
        except (ValueError, ValidationError):
            continue  # Report invalid business data with its row below.
        if opening:
            require_role("finance")(current_user)

    imported = 0
    for idx, row in rows:
        try:
            with db.begin_nested():
                code = _excel_text(row.get("项目编号"))
                customer_name = _excel_text(row.get("客户名称"))
                name = _excel_text(row.get("项目名称")) or customer_name
                missing = [label for label, value in (("项目编号", code), ("客户名称", customer_name)) if not value]
                if missing:
                    raise ValueError("必要字段为空：" + "、".join(missing))
                opening = _initial_receipt_request(row, "pending", customer_name)

                customer = db.query(Customer).filter(Customer.name == customer_name).first()
                if not customer:
                    customer = Customer(name=customer_name)
                    db.add(customer)
                    db.flush()

                try:
                    contract_amt = Decimal(_excel_amount(row.get("合同金额")) or "0")
                except ArithmeticError:
                    contract_amt = Decimal("0")

                project = db.query(Project).filter(Project.code == code).first()
                if not project:
                    project = Project(code=code, name=name, customer_id=customer.id,
                                      manager_id=current_user.id, budget=contract_amt)
                    db.add(project)
                    db.flush()
                    db.add(ProjectMember(project_id=project.id, user_id=current_user.id, role_in_project="manager"))
                else:
                    project.name, project.customer_id, project.budget = name, customer.id, contract_amt

                contract = db.query(Contract).filter(Contract.project_id == project.id).first()
                if not contract and contract_amt > 0:
                    db.add(Contract(project_id=project.id, code=f"HT-{code}", name=f"{name}合同",
                                    amount=contract_amt, status="active"))

                if opening:
                    # Generated opening references retain exact-repeat behavior,
                    # including reversed historical receipts; mismatches are errors.
                    previous = db.query(Receipt).filter(Receipt.bank_ref == opening.bank_ref).all()
                    if previous:
                        if len(previous) != 1 or (
                            previous[0].project_id, previous[0].amount, previous[0].payer, previous[0].notes
                        ) != (project.id, opening.amount, opening.payer, opening.notes):
                            raise ValueError("初始已收金额或银行流水号与已导入回款不一致，请核对")
                    else:
                        opening.project_id = project.id
                        _save_excel_receipt(opening, db, current_user, confirmed=True)
                db.flush()
            imported += 1
        except (ValueError, ValidationError, HTTPException, IntegrityError) as exc:
            errors.append(f"第{idx}行: {_excel_error(exc, '初始已收金额')}")
    return imported


def _import_receipt_rows(wb, db: Session, current_user: User, errors: list[str]) -> int:
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        errors.append("未找到回款明细表头")
        return 0
    headers = [_excel_text(h) for h in rows[0]]
    detail = [(idx, dict(zip(headers, row))) for idx, row in enumerate(rows[1:], start=2)
              if any(_excel_text(value) for value in row)]
    for _, row in detail:
        code = _excel_text(row.get("项目编号"))
        target = db.query(Project).filter(Project.code == code).first()
        if target:
            require_project(db, current_user, target.id)
    if not detail and not all(label in headers for label in ("项目编号", "付款方", "回款金额", "回款日期")):
        errors.append("回款明细表头缺少项目编号、付款方、回款金额或回款日期列")

    imported = 0
    for idx, row in detail:
        try:
            with db.begin_nested():
                missing = [label for label in ("项目编号", "付款方", "回款金额", "回款日期")
                           if not _excel_text(row.get(label))]
                if missing:
                    raise ValueError("必要字段为空：" + "、".join(missing))
                code = _excel_text(row["项目编号"])
                project = db.query(Project).filter(Project.code == code).first()
                if not project:
                    raise ValueError(f"项目编号 {code} 不存在")
                req = ReceiptCreate(
                    project_id=project.id, amount=_excel_amount(row["回款金额"]),
                    payer=_excel_text(row["付款方"]), received_date=_excel_received_date(row["回款日期"]),
                    bank_ref=_excel_text(row.get("银行流水号")),
                )
                _save_excel_receipt(req, db, current_user)
            imported += 1
        except (ValueError, ValidationError, HTTPException, IntegrityError) as exc:
            errors.append(f"第{idx}行: {_excel_error(exc, '回款金额')}")
    return imported


@import_router.post("/excel")
async def import_excel(file: UploadFile = File(...), import_type: str = Query("projects"), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Excel 导入：权限预检整批拒绝，业务错误逐行报告，有效行一次提交。"""
    require_role("member")(current_user)
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="需要安装 openpyxl")

    contents = await file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
    errors = []
    imported = 0
    try:
        # Once per batch, before any writes or savepoints. Calling this helper
        # per row would roll back earlier successful rows.
        _lock_receipt_write(db)
        if import_type == "projects":
            imported = _import_project_rows(wb, db, current_user, errors)
        elif import_type == "receipts":
            imported = _import_receipt_rows(wb, db, current_user, errors)
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        wb.close()
    return {"imported": imported, "errors": errors, "report": f"成功导入 {imported} 条，{len(errors)} 条错误"}


workbench_router = APIRouter(prefix="/api/workbench", tags=["工作台"])


@workbench_router.get("")
def get_workbench(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project_ids = visible_project_ids(current_user)
    now = datetime.utcnow()
    my_tasks = scope_project_rows(db.query(Task), Task.project_id, current_user).filter(Task.assignee_id == current_user.id, Task.status != "done").order_by(Task.due_date.asc()).limit(20).all()
    due_today = [t for t in my_tasks if t.due_date and t.due_date.date() == now.date()]
    overdue = [t for t in my_tasks if t.due_date and t.due_date < now]
    unread_notifications = db.query(Notification).filter(Notification.user_id == current_user.id, Notification.is_read == False).count()
    my_projects = scope_project_rows(db.query(Project), Project.id, current_user).limit(10).all()
    # 最近活动
    recent = scope_project_rows(db.query(ActivityLog), ActivityLog.project_id, current_user).order_by(ActivityLog.created_at.desc()).limit(10).all()
    return {
        "my_tasks": len(my_tasks),
        "due_today": len(due_today),
        "overdue": len(overdue),
        "unread_notifications": unread_notifications,
        "recent_projects": [{"id": p.id, "name": p.name, "code": p.code} for p in my_projects],
        "recent_activities": [{"id": a.id, "action": a.action, "detail": a.detail, "time": a.created_at.isoformat()} for a in recent],
        "tasks": [{"id": t.id, "title": t.title, "status": t.status, "due_date": t.due_date, "project_id": t.project_id} for t in my_tasks[:10]],
        "task_status_counts": {
            "todo": scope_project_rows(db.query(Task), Task.project_id, current_user).filter(Task.assignee_id == current_user.id, Task.status == "todo").count(),
            "in_progress": scope_project_rows(db.query(Task), Task.project_id, current_user).filter(Task.assignee_id == current_user.id, Task.status == "in_progress").count(),
            "done": scope_project_rows(db.query(Task), Task.project_id, current_user).filter(Task.assignee_id == current_user.id, Task.status == "done").count(),
        },
    }


notification_router = APIRouter(prefix="/api/notifications", tags=["通知"])


@notification_router.get("")
def list_notifications(unread_only: bool = False, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    q = db.query(Notification).filter(Notification.user_id == current_user.id)
    if unread_only:
        q = q.filter(Notification.is_read == False)
    return q.order_by(Notification.created_at.desc()).limit(50).all()


@notification_router.put("/{notification_id}/read")
def mark_read(notification_id: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    n = db.query(Notification).filter(Notification.id == notification_id, Notification.user_id == current_user.id).first()
    if n:
        n.is_read = True
        db.commit()
    return {"ok": True}


@notification_router.delete("/{notification_id}")
def delete_notification(notification_id: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    n = db.query(Notification).filter(Notification.id == notification_id, Notification.user_id == current_user.id).first()
    if not n:
        raise HTTPException(status_code=404, detail="通知不存在")
    db.delete(n)
    db.commit()
    return {"ok": True}


@notification_router.put("/read-all")
def mark_all_read(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    db.query(Notification).filter(Notification.user_id == current_user.id, Notification.is_read == False).update({"is_read": True})
    db.commit()
    return {"ok": True}


audit_router = APIRouter(prefix="/api/audit-logs", tags=["审计日志"])


@audit_router.get("")
def list_audit_logs(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    from sqlalchemy import and_, or_, select
    from app.models import Invoice
    project_ids = visible_project_ids(current_user)
    query = db.query(AuditLog)
    if not is_project_admin(current_user):
        receipt_ids = select(Receipt.id).where(Receipt.project_id.in_(project_ids))
        invoice_ids = select(Invoice.id).join(Contract).where(Contract.project_id.in_(project_ids))
        query = query.filter(or_(
            and_(AuditLog.target_type == "receipt", AuditLog.target_id.in_(receipt_ids)),
            and_(AuditLog.target_type == "invoice", AuditLog.target_id.in_(invoice_ids)),
            and_(AuditLog.target_type == "project", AuditLog.target_id.in_(project_ids)),
            and_(AuditLog.target_type.notin_(["receipt", "invoice", "project"]), AuditLog.user_id == current_user.id),
        ))
    return query.order_by(AuditLog.created_at.desc()).limit(100).all()


activity_router = APIRouter(prefix="/api/activities", tags=["活动日志"])


@activity_router.get("")
def list_activities(project_id: Optional[str] = Query(None), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    q = scope_project_rows(db.query(ActivityLog), ActivityLog.project_id, current_user).order_by(ActivityLog.created_at.desc())
    if project_id:
        require_project(db, current_user, project_id)
        q = q.filter(ActivityLog.project_id == project_id)
    return q.limit(100).all()
