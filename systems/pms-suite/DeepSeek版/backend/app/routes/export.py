"""Excel 导出路由"""
from io import BytesIO
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
from datetime import datetime
from app.database import get_db
from app.models import Project, Task, Customer, Contract, Receipt, ReceiptAllocation, User
from app.auth import get_current_user
from app.routes.project_access import require_project, scope_project_rows, visible_project_ids
from sqlalchemy import func

router = APIRouter(prefix="/api/export", tags=["导出"])


@router.get("/projects")
def export_projects(
    search: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = scope_project_rows(db.query(Project), Project.id, current_user).filter(Project.is_archived == False)
    if search:
        q = q.filter(Project.name.contains(search) | Project.code.contains(search))
    if status:
        q = q.filter(Project.status == status)
    projects = q.order_by(Project.created_at.desc()).limit(500).all()

    if not projects:
        from fastapi.responses import Response
        return Response(content="No data", media_type="text/plain")

    # Batch pre-load: users, receipts, contracts (eliminate N+1)
    user_ids = {p.manager_id for p in projects}
    users = {u.id: u for u in db.query(User).filter(User.id.in_(user_ids)).all()}
    project_ids = [p.id for p in projects]
    receipts_map = {}
    for row in db.query(Receipt.project_id, func.sum(Receipt.amount)).filter(
        Receipt.project_id.in_(project_ids), Receipt.status == "confirmed"
    ).group_by(Receipt.project_id).all():
        receipts_map[row[0]] = row[1] or 0
    contracts_map = {}
    for row in db.query(Contract.project_id, func.sum(Contract.amount)).filter(
        Contract.project_id.in_(project_ids), Contract.status == "active"
    ).group_by(Contract.project_id).all():
        contracts_map[row[0]] = row[1] or 0

    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="需要安装 openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "项目列表"
    ws.append(["项目编号", "项目名称", "客户", "状态", "负责人", "合同金额", "已收金额", "应收金额", "回款率", "进度"])

    for p in projects:
        customer_name = p.customer.name if p.customer else ""
        manager = users.get(p.manager_id)
        manager_name = manager.real_name if manager else ""
        total_received = receipts_map.get(p.id, 0)
        contract_total = contracts_map.get(p.id, 0)
        receivable = contract_total - total_received
        rate = round(float(total_received / contract_total * 100), 1) if contract_total else 0

        ws.append([
            p.code, p.name, customer_name, p.status,
            manager_name, float(contract_total), float(total_received),
            float(receivable), f"{rate}%", p.progress,
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=projects.xlsx"})


@router.get("/receipts")
def export_receipts(
    project_id: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = scope_project_rows(db.query(Receipt), Receipt.project_id, current_user)
    if project_id:
        require_project(db, current_user, project_id)
        q = q.filter(Receipt.project_id == project_id)
    receipts = q.order_by(Receipt.received_date.desc()).limit(500).all()

    if not receipts:
        from fastapi.responses import Response
        return Response(content="No data", media_type="text/plain")

    # Batch pre-load projects
    proj_ids = {r.project_id for r in receipts}
    projects_map = {p.id: p for p in db.query(Project).filter(Project.id.in_(proj_ids)).all()}

    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="需要安装 openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "回款记录"
    ws.append(["项目", "金额", "付款方", "回款日期", "银行流水号", "支付方式", "状态", "确认时间"])

    for r in receipts:
        proj = projects_map.get(r.project_id)
        proj_name = proj.name if proj else ""
        ws.append([
            proj_name, float(r.amount), r.payer,
            r.received_date.strftime("%Y-%m-%d") if r.received_date else "",
            r.bank_ref or "", r.payment_method, r.status,
            r.confirmed_at.strftime("%Y-%m-%d %H:%M") if r.confirmed_at else "",
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=receipts.xlsx"})


@router.get("/aging")
def export_aging(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project_ids = visible_project_ids(current_user)
    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "回款账龄"
    ws.append(["项目编号", "项目名称", "客户名称", "回款金额", "回款日期", "账龄(天)", "状态"])
    receipts = db.query(Receipt).join(Project, Receipt.project_id == Project.id)\
        .join(Customer, Project.customer_id == Customer.id)\
        .filter(Receipt.status.in_(["draft", "confirmed"]), Receipt.project_id.in_(project_ids))\
        .order_by(Receipt.received_date.desc()).all()
    for r in receipts:
        aging = (datetime.utcnow() - r.received_date).days if r.received_date else 0
        ws.append([
            r.project.code, r.project.name, r.project.customer.name,
            float(r.amount), r.received_date.strftime("%Y-%m-%d") if r.received_date else "",
            aging, r.status,
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=receivable_aging.xlsx"})


@router.get("/contracts")
def export_contracts(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    project_ids = visible_project_ids(current_user)
    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合同列表"
    ws.append(["合同编号", "合同名称", "项目名称", "客户名称", "金额", "签订日期", "状态"])
    contracts = db.query(Contract).join(Project, Contract.project_id == Project.id)\
        .outerjoin(Customer, Project.customer_id == Customer.id)\
        .filter(Contract.project_id.in_(project_ids))\
        .order_by(Contract.signed_date.desc()).all()
    for c in contracts:
        ws.append([c.code, c.name, c.project.name, c.project.customer.name if c.project.customer else "",
                   float(c.amount), c.signed_date.strftime("%Y-%m-%d") if c.signed_date else "", c.status])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=contracts.xlsx"})


@router.get("/customers")
def export_customers(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    try:
        import openpyxl
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "客户列表"
    ws.append(["客户名称", "简称", "信用等级", "行业", "联系人", "联系电话", "状态"])
    for c in db.query(Customer).order_by(Customer.created_at.desc()).all():
        ws.append([c.name, c.short_name or "", c.credit_level or "", c.industry or "",
                   c.contact_person or "", c.contact_phone or "", "active" if c.is_active else "inactive"])
    output = BytesIO(); wb.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=customers.xlsx"})
