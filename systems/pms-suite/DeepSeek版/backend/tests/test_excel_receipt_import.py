"""Synthetic XLSX requests exercise the real import route and disposable SQLite."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from uuid import uuid4

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import event, select
from sqlalchemy.exc import IntegrityError

from app.auth import create_access_token
from app.database import Base, SessionLocal
from app.main import app
from app.models import ActivityLog, AuditLog, Customer, Project, ProjectMember, Receipt, User


RECEIPT_HEADERS = ["项目编号", "付款方", "回款金额", "回款日期", "银行流水号"]
PROJECT_HEADERS = ["项目编号", "项目名称", "客户名称", "合同金额", "初始已收金额"]


@pytest.fixture
def import_case():
    suffix = uuid4().hex
    ids = {name: f"{name[:3]}{uuid4().hex[:24]}" for name in ("admin", "owner", "member", "finance", "outsider")}
    roles = {"admin": "super_admin", "owner": "pm", "member": "member", "finance": "finance", "outsider": "member"}
    code, other_code = f"EXCEL-{suffix}", f"OTHER-{suffix}"
    with SessionLocal() as db:
        for name, uid in ids.items():
            db.add(User(id=uid, username=f"{name}-{suffix}", hashed_password="synthetic-unusable",
                        real_name=f"合成{name}", role=roles[name]))
        customer = Customer(name=f"合成导入客户-{suffix}")
        db.add(customer)
        db.flush()
        project = Project(code=code, name="合成导入项目", customer_id=customer.id, manager_id=ids["owner"])
        other = Project(code=other_code, name="合成其他项目", customer_id=customer.id, manager_id=ids["outsider"])
        db.add_all([project, other])
        db.flush()
        for name in ("member", "finance"):
            db.add(ProjectMember(project_id=project.id, user_id=ids[name]))
        db.commit()
        project_id, other_id = project.id, other.id
    client = TestClient(app, client=(f"excel-{suffix}", 50000))
    yield {"client": client, "ids": ids, "code": code, "project_id": project_id,
           "other_code": other_code, "other_id": other_id,
           "headers": {name: {"Authorization": f"Bearer {create_access_token({'sub': uid})}"}
                       for name, uid in ids.items()}, "suffix": suffix}
    client.close()


def upload(case, rows, *, kind="receipts", actor="admin", columns=None):
    workbook = openpyxl.Workbook()
    workbook.active.append(columns if columns is not None else (RECEIPT_HEADERS if kind == "receipts" else PROJECT_HEADERS))
    for row in rows:
        workbook.active.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return case["client"].post(
        f"/api/import/excel?import_type={kind}", headers=case["headers"][actor],
        files={"file": ("synthetic.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def receipt_row(case, amount="25.25", bank_ref="", **changes):
    row = {"项目编号": case["code"], "付款方": "合成付款方", "回款金额": amount,
           "回款日期": "2026-10-28", "银行流水号": bank_ref}
    row.update(changes)
    return [row[key] for key in RECEIPT_HEADERS]


def database_snapshot():
    with SessionLocal() as db:
        return {name: tuple(tuple(row) for row in db.execute(select(table).order_by(*table.primary_key.columns)))
                for name, table in Base.metadata.tables.items()}


INVALID_AMOUNTS = ["-1.00", "0", "NaN", "sNaN", "Infinity", "-Infinity", "1.001",
                   "0.0001", "10000000000000000", "9999999999999999999", "bad", "¥bad"]


@pytest.mark.parametrize("amount", INVALID_AMOUNTS)
def test_invalid_receipt_amount_leaves_no_rows_or_audit(import_case, amount):
    before = database_snapshot()
    response = upload(import_case, [receipt_row(import_case, amount)])
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0
    assert len(response.json()["errors"]) == 1
    assert "第2行" in response.json()["errors"][0] and "回款金额" in response.json()["errors"][0]
    assert database_snapshot() == before


def test_existing_bank_reference_is_not_imported_twice(import_case):
    case = import_case
    bank_ref = f"SYNTH-{case['suffix']}"
    created = case["client"].post("/api/receipts", headers=case["headers"]["admin"], json={
        "project_id": case["project_id"], "amount": "25.25", "payer": "合成付款方",
        "received_date": "2026-10-28", "bank_ref": bank_ref,
    })
    assert created.status_code == 200, created.text
    before = database_snapshot()
    response = upload(case, [receipt_row(case, bank_ref=bank_ref)])
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0
    assert "银行流水号已存在" in response.json()["errors"][0]
    assert database_snapshot() == before


@pytest.mark.parametrize("label", RECEIPT_HEADERS[:-1])
@pytest.mark.parametrize("empty", [None, "   "])
def test_required_receipt_fields_are_named_per_row(import_case, label, empty):
    before = database_snapshot()
    response = upload(import_case, [receipt_row(import_case, **{label: empty})])
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0
    assert len(response.json()["errors"]) == 1
    assert f"第2行" in response.json()["errors"][0]
    assert label in response.json()["errors"][0]
    assert database_snapshot() == before


@pytest.mark.parametrize("value", ["not-a-date", "2026-02-30", "17"])
def test_invalid_date_is_not_replaced_with_today(import_case, value):
    before = database_snapshot()
    response = upload(import_case, [receipt_row(import_case, **{"回款日期": value})])
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0
    assert "第2行" in response.json()["errors"][0] and "回款日期" in response.json()["errors"][0]
    assert database_snapshot() == before


@pytest.mark.parametrize("amount,date_value,expected", [
    ("¥ 1,234.50", "2026-10-28", "1234.50"),
    (0.01, datetime(2026, 10, 28, 14, 15, 16), "0.01"),
    ("25.00", "2026-10-28T10:11:12", "25.00"),
    ("9999999999999.99", "2026-10-28", "9999999999999.99"),
])
def test_valid_excel_money_dates_and_audit_are_saved(import_case, amount, date_value, expected):
    case = import_case
    response = upload(case, [receipt_row(case, amount, **{"回款日期": date_value})], actor="member")
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1 and response.json()["errors"] == []
    expected_date = date_value if isinstance(date_value, datetime) else datetime.fromisoformat(date_value)
    with SessionLocal() as db:
        receipt = db.query(Receipt).filter_by(project_id=case["project_id"]).one()
        assert (receipt.amount, receipt.received_date, receipt.bank_ref, receipt.status) == (
            Decimal(expected), expected_date, "", "draft")
        audit = db.query(AuditLog).filter_by(target_id=receipt.id).one()
        assert (audit.action, Decimal(audit.new_value), audit.user_id) == ("create", Decimal(expected), case["ids"]["member"])
        assert db.query(ActivityLog).filter_by(user_id=case["ids"]["member"], action="登记回款").count() == 1


def test_valid_rows_survive_invalid_rows_and_duplicate_in_same_batch(import_case):
    case = import_case
    response = upload(case, [
        receipt_row(case, "10.00", "SAME"),
        receipt_row(case, "2.001", "INVALID"),
        receipt_row(case, "20.00", " SAME "),
        receipt_row(case, "30.00", None),
        receipt_row(case, "40.00", None),
    ])
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["imported"] == 3 and len(data["errors"]) == 2
    assert "第3行" in data["errors"][0] and "回款金额" in data["errors"][0]
    assert "第4行" in data["errors"][1] and "银行流水号已存在" in data["errors"][1]
    with SessionLocal() as db:
        receipts = db.query(Receipt).filter_by(project_id=case["project_id"]).all()
        assert sorted(r.amount for r in receipts) == [Decimal("10.00"), Decimal("30.00"), Decimal("40.00")]
        assert sorted(r.bank_ref for r in receipts) == ["", "", "SAME"]
        assert db.query(AuditLog).filter_by(user_id=case["ids"]["admin"], action="create").count() == 3


def test_reversed_bank_reference_can_be_reused_once(import_case):
    case = import_case
    with SessionLocal() as db:
        old = Receipt(project_id=case["project_id"], amount=Decimal("10"), received_date=datetime(2026, 10, 1),
                      bank_ref="REVERSED", status="reversed")
        db.add(old)
        db.commit()
    response = upload(case, [receipt_row(case, bank_ref="REVERSED"), receipt_row(case, bank_ref="REVERSED")])
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1 and len(response.json()["errors"]) == 1
    with SessionLocal() as db:
        assert db.query(Receipt).filter_by(project_id=case["project_id"], bank_ref="REVERSED").count() == 2


@pytest.mark.parametrize("kind", ["receipts", "projects"])
def test_schema_limit_that_sqlite_cannot_store_exactly_is_rejected(import_case, kind):
    case = import_case
    before = database_snapshot()
    if kind == "receipts":
        rows = [receipt_row(case, "9999999999999999.99")]
    else:
        rows = [[f"NEW-{case['suffix']}", "不应新增项目", f"不应新增客户-{case['suffix']}", "100", "9999999999999999.99"]]
    response = upload(case, rows, kind=kind)
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0
    assert "精确保存" in response.json()["errors"][0]
    assert database_snapshot() == before


@pytest.mark.parametrize("amount", [value for value in INVALID_AMOUNTS if value != "0"])
def test_invalid_opening_receipt_does_not_leave_customer_project_contract_or_audit(import_case, amount):
    case = import_case
    before = database_snapshot()
    response = upload(case, [[f"NEW-{case['suffix']}", "合成新项目", f"合成新客户-{case['suffix']}", "100", amount]], kind="projects")
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 0 and len(response.json()["errors"]) == 1
    assert "第2行" in response.json()["errors"][0] and "初始已收金额" in response.json()["errors"][0]
    assert database_snapshot() == before


@pytest.mark.parametrize("amount", [None, "0", "0.00"])
def test_blank_or_zero_opening_amount_still_imports_project_without_receipt(import_case, amount):
    case = import_case
    code = f"NEW-{case['suffix']}"
    response = upload(case, [[code, "合成新项目", f"合成新客户-{case['suffix']}", "100", amount]], kind="projects", actor="member")
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1 and response.json()["errors"] == []
    with SessionLocal() as db:
        project = db.query(Project).filter_by(code=code).one()
        assert db.query(Receipt).filter_by(project_id=project.id).count() == 0
        assert db.query(AuditLog).filter_by(user_id=case["ids"]["member"], target_type="receipt").count() == 0


def test_opening_receipt_repeat_is_exact_and_audited_once(import_case):
    case = import_case
    code = f"NEW-{case['suffix']}"
    rows = [[code, "合成新项目", f"合成新客户-{case['suffix']}", "100", "25.25"]]
    first = upload(case, rows, kind="projects", actor="finance")
    assert first.status_code == 200 and first.json()["imported"] == 1, first.text
    repeated = upload(case, rows, kind="projects", actor="finance")
    assert repeated.status_code == 200 and repeated.json()["imported"] == 1, repeated.text
    with SessionLocal() as db:
        project = db.query(Project).filter_by(code=code).one()
        receipt = db.query(Receipt).filter_by(project_id=project.id).one()
        assert (receipt.amount, receipt.bank_ref, receipt.status, receipt.confirmed_by) == (
            Decimal("25.25"), f"IMPORT-{code}", "confirmed", case["ids"]["finance"])
        assert receipt.confirmed_at is not None
        assert sorted(a.action for a in db.query(AuditLog).filter_by(target_id=receipt.id)) == ["confirm", "create"]
        assert db.query(ActivityLog).filter_by(user_id=case["ids"]["finance"], action="登记回款").count() == 1
    before = database_snapshot()
    mismatch = upload(case, [[code, "不应更新名称", "不应新增客户", "999", "30.00"]], kind="projects", actor="finance")
    assert mismatch.status_code == 200 and mismatch.json()["imported"] == 0, mismatch.text
    assert "不一致" in mismatch.json()["errors"][0]
    assert database_snapshot() == before
    reversed_response = case["client"].put(f"/api/receipts/{receipt.id}/reverse", headers=case["headers"]["finance"])
    assert reversed_response.status_code == 200, reversed_response.text
    before = database_snapshot()
    repeated = upload(case, rows, kind="projects", actor="finance")
    assert repeated.status_code == 200 and repeated.json()["imported"] == 1, repeated.text
    assert database_snapshot() == before  # An import replay must not resurrect a reversed receipt.


@pytest.mark.parametrize("same_project", [False, True])
def test_opening_bank_collision_rolls_back_existing_project_changes(import_case, same_project):
    case = import_case
    with SessionLocal() as db:
        db.add(Receipt(project_id=case["project_id"] if same_project else case["other_id"],
                       amount=Decimal("25.25"), received_date=datetime(2026, 10, 28),
                       bank_ref=f"IMPORT-{case['code']}", status="draft"))
        db.commit()
    before = database_snapshot()
    response = upload(case, [[case["code"], "不应更新项目", "不应新增客户", "200", "25.25"]], kind="projects")
    assert response.status_code == 200 and response.json()["imported"] == 0, response.text
    assert "银行流水号" in response.json()["errors"][0]
    assert database_snapshot() == before


@pytest.mark.parametrize("kind", ["receipts", "projects"])
def test_late_audit_failure_rolls_back_only_that_row(import_case, kind):
    case = import_case
    attempts = 0

    def reject_second_audit(mapper, connection, target):
        nonlocal attempts
        if target.user_id == case["ids"]["admin"] and target.action == "create":
            attempts += 1
            if attempts == 2:
                raise IntegrityError("synthetic audit failure", {}, RuntimeError("synthetic constraint"))

    if kind == "receipts":
        rows = [receipt_row(case, amount, f"LATE-{i}") for i, amount in enumerate(["10", "20", "30"])]
    else:
        rows = [[f"NEW-{case['suffix']}-{i}", f"合成项目{i}", f"合成客户-{case['suffix']}-{i}", "100", amount]
                for i, amount in enumerate(["10", "20", "30"])]
    event.listen(AuditLog, "before_insert", reject_second_audit)
    try:
        response = upload(case, rows, kind=kind)
    finally:
        event.remove(AuditLog, "before_insert", reject_second_audit)
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 2 and len(response.json()["errors"]) == 1
    assert "第3行" in response.json()["errors"][0]
    assert attempts == 3
    with SessionLocal() as db:
        logs = db.query(AuditLog).filter_by(user_id=case["ids"]["admin"], action="create").all()
        assert sorted(Decimal(a.new_value) for a in logs) == [Decimal("10"), Decimal("30")]
        receipts = db.query(Receipt).filter(Receipt.id.in_([a.target_id for a in logs])).all()
        assert sorted(r.amount for r in receipts) == [Decimal("10"), Decimal("30")]
        assert db.query(ActivityLog).filter_by(user_id=case["ids"]["admin"], action="登记回款").count() == 2
        if kind == "receipts":
            assert db.query(Receipt).filter_by(project_id=case["project_id"]).count() == 2
        else:
            assert db.query(Customer).filter_by(name=f"合成客户-{case['suffix']}-1").count() == 0
            assert db.query(Project).filter_by(code=f"NEW-{case['suffix']}-1").count() == 0
            from app.models import Contract
            assert db.query(Contract).filter_by(code=f"HT-NEW-{case['suffix']}-1").count() == 0


@pytest.mark.parametrize("kind", ["receipts", "projects"])
def test_commit_failure_does_not_commit_released_savepoints(import_case, kind, monkeypatch):
    from sqlalchemy.orm import Session

    case = import_case
    before = database_snapshot()
    original_commit = Session.commit

    def reject_commit(self):
        if self.get_bind().dialect.name == "sqlite":
            raise RuntimeError("synthetic commit failure")
        return original_commit(self)

    rows = ([receipt_row(case)] if kind == "receipts" else
            [[f"NEW-{case['suffix']}", "不应新增项目", f"不应新增客户-{case['suffix']}", "100", "25.25"]])
    monkeypatch.setattr(Session, "commit", reject_commit)
    with pytest.raises(RuntimeError, match="synthetic commit failure"):
        upload(case, rows, kind=kind)
    assert database_snapshot() == before


@pytest.mark.parametrize("kind", ["receipts", "projects"])
def test_foreign_project_preflight_rejects_whole_batch_without_writes(import_case, kind):
    case = import_case
    if kind == "receipts":
        rows = [receipt_row(case), receipt_row(case, **{"项目编号": case["other_code"]})]
    else:
        rows = [[code, "不应更新项目", "不应新增客户", "100", "0"] for code in [case["code"], case["other_code"]]]
    before = database_snapshot()
    response = upload(case, rows, kind=kind, actor="owner")
    assert response.status_code == 404, response.text
    assert database_snapshot() == before


def test_opening_confirmation_keeps_existing_finance_permission(import_case):
    case = import_case
    before = database_snapshot()
    response = upload(case, [[f"NEW-{case['suffix']}", "不应新增项目", "不应新增客户", "100", "25.25"]],
                      kind="projects", actor="member")
    assert response.status_code == 403, response.text
    assert database_snapshot() == before


def test_concurrent_excel_and_json_bank_reference_has_one_receipt_and_audit(import_case):
    case = import_case
    bank_ref = f"RACE-{case['suffix']}"
    barrier = Barrier(2)

    def post_excel():
        barrier.wait(timeout=10)
        return upload(case, [receipt_row(case, bank_ref=bank_ref)])

    def post_json():
        barrier.wait(timeout=10)
        with TestClient(app, client=(f"excel-json-{case['suffix']}", 50000)) as client:
            return client.post("/api/receipts", headers=case["headers"]["admin"], json={
                "project_id": case["project_id"], "amount": "25.25", "payer": "合成付款方",
                "received_date": "2026-10-28", "bank_ref": bank_ref,
            })

    with ThreadPoolExecutor(max_workers=2) as pool:
        excel_future, json_future = pool.submit(post_excel), pool.submit(post_json)
        excel, json_response = excel_future.result(timeout=20), json_future.result(timeout=20)
    assert excel.status_code == 200, excel.text
    assert json_response.status_code in (200, 400), json_response.text
    assert excel.json()["imported"] + (json_response.status_code == 200) == 1
    with SessionLocal() as db:
        receipt = db.query(Receipt).filter_by(project_id=case["project_id"], bank_ref=bank_ref).one()
        assert receipt.amount == Decimal("25.25")
        assert db.query(AuditLog).filter_by(target_id=receipt.id, action="create").count() == 1


def test_missing_columns_report_all_missing_business_fields_on_the_row(import_case):
    before = database_snapshot()
    response = upload(import_case, [[import_case["code"], "合成付款方"]], columns=["项目编号", "付款方"])
    assert response.status_code == 200 and response.json()["imported"] == 0, response.text
    errors = response.json()["errors"]
    assert len(errors) == 1 and "第2行" in errors[0] and "回款金额" in errors[0] and "回款日期" in errors[0]
    assert database_snapshot() == before


def test_unknown_project_is_reported_without_creating_any_rows(import_case):
    before = database_snapshot()
    response = upload(import_case, [receipt_row(import_case, **{"项目编号": "SYNTHETIC-UNKNOWN"})])
    assert response.status_code == 200 and response.json()["imported"] == 0, response.text
    assert "第2行" in response.json()["errors"][0] and "不存在" in response.json()["errors"][0]
    assert database_snapshot() == before


def test_empty_workbook_has_a_report_instead_of_an_index_error(import_case):
    workbook = openpyxl.Workbook()
    stream = BytesIO()
    workbook.save(stream)
    before = database_snapshot()
    response = import_case["client"].post(
        "/api/import/excel?import_type=receipts", headers=import_case["headers"]["admin"],
        files={"file": ("synthetic-empty.xlsx", stream.getvalue())},
    )
    assert response.status_code == 200 and response.json()["imported"] == 0, response.text
    assert response.json()["errors"] == ["未找到回款明细表头"]
    assert database_snapshot() == before


def test_project_sheet_selection_preserves_excel_row_numbers_and_skips_totals(import_case):
    case = import_case
    workbook = openpyxl.Workbook()
    workbook.active.append(["合成说明页"])
    detail = workbook.create_sheet("合成明细")
    detail.append(["合成说明"])
    detail.append([])
    detail.append(PROJECT_HEADERS)
    code = f"NEW-{case['suffix']}"
    detail.append([code, "合成项目", f"合成客户-{case['suffix']}", "100", "10.25"])
    detail.append([f"BAD-{case['suffix']}", "不应新增项目", "不应新增客户", "100", "NaN"])
    detail.append(["合计", "总计", None, "200", "NaN"])
    stream = BytesIO()
    workbook.save(stream)
    response = case["client"].post(
        "/api/import/excel?import_type=projects", headers=case["headers"]["admin"],
        files={"file": ("synthetic-sheets.xlsx", stream.getvalue())},
    )
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1 and len(response.json()["errors"]) == 1
    assert "第5行" in response.json()["errors"][0] and "初始已收金额" in response.json()["errors"][0]
    with SessionLocal() as db:
        assert db.query(Receipt).filter_by(bank_ref=f"IMPORT-{code}").one().amount == Decimal("10.25")
        assert db.query(Project).filter_by(code=f"BAD-{case['suffix']}").count() == 0
        assert db.query(Customer).filter_by(name="不应新增客户").count() == 0
