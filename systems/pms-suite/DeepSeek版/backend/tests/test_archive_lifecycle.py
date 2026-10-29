"""WP-27: exercise real archive APIs against conftest's disposable SQLite DB."""
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from threading import Barrier
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.database import SessionLocal
from app.main import app
from app.models import AuditLog, Receipt, ReceiptAllocation, User


@pytest.fixture
def archive_case():
    # Each independent caller has its own IP; do not disable the rate limiter.
    with TestClient(app, client=(f"archive-{uuid4().hex}", 50000)) as client:
        login = client.post("/api/auth/login", json={
            "username": "test_admin", "password": "Admin@123456",
        })
        assert login.status_code == 200
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        suffix = uuid4().hex
        customer = client.post("/api/customers", headers=headers, json={"name": f"虚构归档客户-{suffix}"})
        assert customer.status_code == 200
        project = client.post("/api/projects", headers=headers, json={
            "code": f"ARCH-{suffix}", "name": "虚构关务实施项目",
            "customer_id": customer.json()["id"], "manager_id": "test_admin",
        })
        assert project.status_code == 200
        contract = client.post("/api/contracts", headers=headers, json={
            "code": f"CT-{suffix}", "name": "虚构实施合同", "amount": "100.00",
            "project_id": project.json()["id"],
        })
        assert contract.status_code == 200
        invoice = client.post("/api/invoices", headers=headers, json={
            "code": f"INV-{suffix}", "contract_id": contract.json()["id"], "amount": "100.00",
        })
        assert invoice.status_code == 200
        yield client, headers, project.json()["id"], invoice.json()["id"], customer.json()["id"]


def receipt_body(project_id, **changes):
    return {
        "project_id": project_id, "amount": "25.00",
        "received_date": "2026-01-15T00:00:00", "payer": "虚构付款方",
        **changes,
    }


def receipt_count(project_id):
    with SessionLocal() as db:
        return db.query(Receipt).filter(Receipt.project_id == project_id).count()


def test_bad_password_does_not_issue_token_or_change_account(archive_case):
    client, headers, *_ = archive_case
    with SessionLocal() as db:
        before = db.query(User).filter_by(id="test_admin").one().hashed_password
    response = client.post("/api/auth/login", json={"username": "test_admin", "password": "wrong"})
    assert response.status_code == 401
    assert "access_token" not in response.json()
    assert client.get("/api/auth/me", headers=headers).json()["id"] == "test_admin"
    with SessionLocal() as db:
        assert db.query(User).filter_by(id="test_admin").one().hashed_password == before


def test_receipt_lifecycle_retains_money_and_audit(archive_case):
    client, headers, project_id, invoice_id, _ = archive_case
    response = client.post("/api/receipts", headers=headers, json=receipt_body(
        project_id, allocations=[{"invoice_id": invoice_id, "allocated_amount": "25.00"}],
    ))
    assert response.status_code == 200
    receipt_id = response.json()["id"]
    stats_url = f"/api/receipts/project/{project_id}/stats"
    assert client.get(stats_url, headers=headers).json()["total_received"] == 0
    assert client.put(f"/api/receipts/{receipt_id}/confirm", headers=headers).status_code == 200
    stats = client.get(stats_url, headers=headers).json()
    # Invoice registration currently creates a draft; no issue transition exists.
    assert stats == {"total_contract": 100, "total_received": 25, "total_invoiced": 0,
                     "receivable": 75, "receipt_rate": 25}
    assert client.put(f"/api/receipts/{receipt_id}/confirm", headers=headers).status_code == 400
    assert client.put(f"/api/receipts/{receipt_id}/reverse", headers=headers).status_code == 200
    assert client.put(f"/api/receipts/{receipt_id}/reverse", headers=headers).status_code == 400
    assert client.get(stats_url, headers=headers).json()["total_received"] == 0
    with SessionLocal() as db:
        row = db.get(Receipt, receipt_id)
        assert (row.amount, row.status, row.version) == (Decimal("25.00"), "reversed", 2)
        assert db.query(ReceiptAllocation).filter_by(receipt_id=receipt_id).one().allocated_amount == Decimal("25.00")
        audit = db.query(AuditLog).filter_by(target_type="receipt", target_id=receipt_id).all()
        assert sorted(a.action for a in audit) == ["confirm", "create", "reverse"]


def test_idempotency_replay_preserves_one_receipt_and_one_audit(archive_case):
    client, headers, project_id, _, _ = archive_case
    body = receipt_body(project_id, idempotency_key=uuid4().hex)
    first = client.post("/api/finance/receipts", headers=headers, json=body)
    replay = client.post("/api/finance/receipts", headers=headers, json=body)
    assert first.status_code == 201
    assert replay.status_code == 200
    assert first.json()["id"] == replay.json()["id"]
    assert receipt_count(project_id) == 1
    with SessionLocal() as db:
        assert db.query(AuditLog).filter_by(target_type="receipt", target_id=first.json()["id"], action="create").count() == 1


@pytest.mark.parametrize("keys", [("order-10", "order-1"), ("wildcard-real", "wildcard-%")])
def test_distinct_literal_idempotency_keys_never_match_substrings(archive_case, keys):
    client, headers, project_id, _, _ = archive_case
    suffix = uuid4().hex
    ids = []
    for key in keys:
        response = client.post("/api/finance/receipts", headers=headers, json=receipt_body(
            project_id, idempotency_key=f"{suffix}-{key}",
        ))
        assert response.status_code == 201
        ids.append(response.json()["id"])
    assert len(set(ids)) == 2
    assert receipt_count(project_id) == 2


def test_reused_key_with_changed_payload_is_conflict_and_no_write(archive_case):
    client, headers, project_id, _, _ = archive_case
    body = receipt_body(project_id, idempotency_key=uuid4().hex)
    assert client.post("/api/finance/receipts", headers=headers, json=body).status_code == 201
    conflict = client.post("/api/finance/receipts", headers=headers, json={**body, "amount": "30.00"})
    assert conflict.status_code == 409
    assert receipt_count(project_id) == 1
    with SessionLocal() as db:
        assert db.query(Receipt).filter_by(project_id=project_id).one().amount == Decimal("25.00")


def test_concurrent_replay_commits_only_once(archive_case):
    _, headers, project_id, _, _ = archive_case
    body = receipt_body(project_id, idempotency_key=uuid4().hex)
    barrier = Barrier(4)

    def submit(_):
        with TestClient(app, client=(f"parallel-{uuid4().hex}", 50000)) as client:
            barrier.wait(timeout=10)
            return client.post("/api/finance/receipts", headers=headers, json=body)

    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(submit, range(4)))
    assert sorted(r.status_code for r in responses) == [200, 200, 200, 201]
    assert len({r.json()["id"] for r in responses}) == 1
    assert receipt_count(project_id) == 1


@pytest.mark.parametrize("endpoint", ["/api/receipts", "/api/finance/receipts"])
@pytest.mark.parametrize("amount", ["0", "-1", "1.001"])
def test_invalid_money_is_rejected_without_receipt(archive_case, endpoint, amount):
    client, headers, project_id, _, _ = archive_case
    response = client.post(endpoint, headers=headers, json=receipt_body(project_id, amount=amount))
    assert response.status_code == 422
    assert receipt_count(project_id) == 0


@pytest.mark.parametrize("endpoint", ["/api/receipts", "/api/finance/receipts"])
def test_invalid_allocation_rolls_back_entire_receipt(archive_case, endpoint):
    client, headers, project_id, invoice_id, _ = archive_case
    key = uuid4().hex
    response = client.post(endpoint, headers=headers, json=receipt_body(
        project_id, idempotency_key=key,
        allocations=[{"invoice_id": invoice_id, "allocated_amount": "24.00"}],
    ))
    assert response.status_code == 400
    assert receipt_count(project_id) == 0
    with SessionLocal() as db:
        assert db.query(ReceiptAllocation).filter_by(invoice_id=invoice_id).count() == 0
    # The failed request must not consume a key or prevent a valid retry.
    body = receipt_body(project_id, idempotency_key=key,
                        allocations=[{"invoice_id": invoice_id, "allocated_amount": "25.00"}])
    response = client.post("/api/finance/receipts", headers=headers, json=body)
    assert response.status_code == 201


@pytest.mark.parametrize("endpoint", ["/api/receipts", "/api/finance/receipts"])
def test_allocation_cannot_exceed_invoice_balance(archive_case, endpoint):
    client, headers, project_id, invoice_id, _ = archive_case
    body = receipt_body(project_id, amount="75.00",
                        allocations=[{"invoice_id": invoice_id, "allocated_amount": "75.00"}])
    assert client.post(endpoint, headers=headers, json=body).status_code == (201 if endpoint.startswith('/api/finance') else 200)
    response = client.post(endpoint, headers=headers, json=body)
    assert response.status_code == 400
    assert receipt_count(project_id) == 1


def test_portal_is_limited_to_linked_customer_and_denies_internal_apis(archive_case):
    client, headers, project_id, _, customer_id = archive_case
    username = f"portal-{uuid4().hex}"
    grant = client.post(f"/api/customers/{customer_id}/portal-users", headers=headers,
                        json={"username": username, "password": "Synthetic@123"})
    assert grant.status_code == 200
    wrong = client.post("/api/portal/login", json={"username": username, "password": "wrong"})
    assert wrong.status_code == 401
    login = client.post("/api/portal/login", json={"username": username, "password": "Synthetic@123"})
    assert login.status_code == 200
    portal_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    projects = client.get("/api/portal/projects", headers=portal_headers)
    assert projects.status_code == 200
    assert {p["id"] for p in projects.json()} == {project_id}
    assert client.get(f"/api/portal/projects/{project_id}", headers=portal_headers).json()["project"]["id"] == project_id
    assert client.get("/api/portal/projects/p002", headers=portal_headers).status_code == 404
    assert client.get("/api/projects", headers=portal_headers).status_code == 403
    assert client.get("/api/finance/receipts", headers=portal_headers).status_code == 403
    assert client.get("/api/portal/projects", headers={"Authorization": "Bearer invalid"}).status_code == 401


@pytest.mark.parametrize("endpoint", ["/api/receipts", "/api/finance/receipts"])
@pytest.mark.parametrize("allocation", [
    {"invoice_id": "missing", "allocated_amount": "25.00"},
    {"invoice_id": "inv002", "allocated_amount": "25.00"},  # another project's invoice
])
def test_missing_or_foreign_invoice_never_leaves_rows(archive_case, endpoint, allocation):
    client, headers, project_id, _, _ = archive_case
    response = client.post(endpoint, headers=headers, json=receipt_body(project_id, allocations=[allocation]))
    assert response.status_code == 400
    assert receipt_count(project_id) == 0


@pytest.mark.parametrize("endpoint", ["/api/receipts", "/api/finance/receipts"])
def test_negative_and_duplicate_allocations_are_rejected(archive_case, endpoint):
    client, headers, project_id, invoice_id, _ = archive_case
    for allocations in [
        [{"invoice_id": invoice_id, "allocated_amount": "-25.00"}],
        [{"invoice_id": invoice_id, "allocated_amount": "10.00"},
         {"invoice_id": invoice_id, "allocated_amount": "15.00"}],
    ]:
        response = client.post(endpoint, headers=headers, json=receipt_body(project_id, allocations=allocations))
        assert response.status_code == 422
    assert receipt_count(project_id) == 0


def test_legacy_note_key_is_preserved_and_requires_review(archive_case):
    client, headers, project_id, _, _ = archive_case
    key = f"legacy-{uuid4().hex}"
    legacy = client.post("/api/receipts", headers=headers, json=receipt_body(project_id, notes=f"旧备注 idempotency:{key}"))
    assert legacy.status_code == 200
    response = client.post("/api/finance/receipts", headers=headers,
                           json=receipt_body(project_id, idempotency_key=key))
    assert response.status_code == 409
    assert receipt_count(project_id) == 1
    with SessionLocal() as db:
        assert db.get(Receipt, legacy.json()["id"]).notes == f"旧备注 idempotency:{key}"


def test_replay_after_reverse_does_not_restore_money(archive_case):
    client, headers, project_id, _, _ = archive_case
    body = receipt_body(project_id, idempotency_key=uuid4().hex)
    first = client.post("/api/finance/receipts", headers=headers, json=body)
    assert first.status_code == 201
    receipt_id = first.json()["id"]
    assert client.put(f"/api/receipts/{receipt_id}/confirm", headers=headers).status_code == 200
    assert client.put(f"/api/receipts/{receipt_id}/reverse", headers=headers).status_code == 200
    replay = client.post("/api/finance/receipts", headers=headers, json=body)
    assert replay.status_code == 200
    assert replay.json()["id"] == receipt_id
    assert replay.json()["status"] == "reversed"
    assert receipt_count(project_id) == 1
    assert client.get(f"/api/receipts/project/{project_id}/stats", headers=headers).json()["total_received"] == 0


def test_member_cannot_confirm_or_reverse_canonical_receipt(archive_case):
    client, headers, project_id, _, _ = archive_case
    first = client.post("/api/receipts", headers=headers, json=receipt_body(project_id))
    assert first.status_code == 200
    receipt_id = first.json()["id"]
    login = client.post("/api/auth/login", json={"username": "dev_li", "password": "dev123"})
    assert login.status_code == 200
    member_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert client.put(f"/api/receipts/{receipt_id}/confirm", headers=member_headers).status_code == 403
    assert client.put(f"/api/receipts/{receipt_id}/confirm", headers=headers).status_code == 200
    assert client.put(f"/api/receipts/{receipt_id}/reverse", headers=member_headers).status_code == 403
    with SessionLocal() as db:
        assert db.get(Receipt, receipt_id).status == "confirmed"


def test_reused_key_by_other_actor_is_rejected(archive_case):
    client, headers, project_id, _, _ = archive_case
    body = receipt_body(project_id, idempotency_key=uuid4().hex)
    assert client.post("/api/finance/receipts", headers=headers, json=body).status_code == 201
    login = client.post("/api/auth/login", json={"username": "finance", "password": "finance123"})
    assert login.status_code == 200
    # This is a key/actor conflict test: both actors must first have project access.
    from app.models import ProjectMember
    with SessionLocal() as db:
        db.add(ProjectMember(project_id=project_id, user_id=login.json()['user']['id']))
        db.commit()
    other_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert client.post("/api/finance/receipts", headers=other_headers, json=body).status_code == 409
    assert receipt_count(project_id) == 1


@pytest.fixture
def security_case(archive_case):
    from app.auth import hash_password
    from app.models import Contract, Project, ProjectMember
    client, admin, project_id, invoice_id, customer_id = archive_case
    suffix = uuid4().hex
    secret = f'虚构保密项目-{suffix}'
    identities = {}
    password_hash = hash_password('Synthetic@123')
    with SessionLocal() as db:
        for name, role in [('owner', 'pm'), ('member', 'member'), ('outsider', 'member'), ('enterprise_admin', 'enterprise_admin'),
                           ('finance_member', 'finance'), ('finance_outsider', 'finance'), ('guest', 'guest')]:
            user = User(username=f'{name}-{suffix}', real_name=f'合成{name}', role=role,
                        hashed_password=password_hash)
            db.add(user)
            db.flush()
            identities[name] = {'id': user.id, 'username': user.username}
        project = db.get(Project, project_id)
        project.manager_id = identities['owner']['id']
        project.name = secret
        contract = db.query(Contract).filter_by(project_id=project_id).one()
        contract.name = secret
        contract_id = contract.id
        for name in ['member', 'finance_member', 'guest']:
            db.add(ProjectMember(project_id=project_id, user_id=identities[name]['id']))
        db.commit()
    headers = {'admin': admin}
    for name, identity in identities.items():
        login = client.post('/api/auth/login', json={'username': identity['username'], 'password': 'Synthetic@123'})
        assert login.status_code == 200
        headers[name] = {'Authorization': f"Bearer {login.json()['access_token']}"}
    other = client.post('/api/projects', headers=admin, json={
        'code': f'OTHER-{suffix}', 'name': '虚构其他项目', 'customer_id': customer_id,
        'manager_id': identities['outsider']['id'],
    })
    assert other.status_code == 200
    body = receipt_body(project_id, payer=secret, idempotency_key=f'scope-{suffix}',
                        allocations=[{'invoice_id': invoice_id, 'allocated_amount': '25.00'}])
    receipt = client.post('/api/finance/receipts', headers=admin, json=body)
    assert receipt.status_code == 201
    task = client.post('/api/tasks', headers=admin, json={'project_id': project_id, 'title': secret, 'assignee_id': identities['member']['id']})
    assert task.status_code == 200
    milestone = client.post('/api/milestones', headers=admin, json={'project_id': project_id, 'name': secret})
    assert milestone.status_code == 200
    document = client.post('/api/documents', headers=admin, json={'project_id': project_id, 'name': secret})
    assert document.status_code == 200
    return {'client': client, 'headers': headers, 'users': identities, 'project': project_id,
            'other_project': other.json()['id'], 'invoice': invoice_id, 'contract': contract_id,
            'receipt': receipt.json()['id'], 'receipt_body': body, 'task': task.json()['id'],
            'milestone': milestone.json()['id'], 'document': document.json()['id'], 'secret': secret}


def security_snapshot():
    """Hash persisted rows, so denied operations cannot hide audit/key side effects."""
    import hashlib
    import json
    from sqlalchemy import select
    from app.database import Base
    result = {}
    with SessionLocal() as db:
        for name, table in Base.metadata.tables.items():
            rows = db.execute(select(table).order_by(*table.primary_key.columns)).all()
            payload = json.dumps([list(row) for row in rows], default=str, ensure_ascii=False)
            result[name] = hashlib.sha256(payload.encode()).hexdigest()
    return result


@pytest.mark.parametrize('actor', ['outsider', 'finance_outsider', 'guest'])
def test_project_read_write_denial_has_no_persisted_effects(security_case, actor):
    case = security_case
    client, headers = case['client'], case['headers'][actor]
    expected = 403 if actor == 'guest' else 404
    before = security_snapshot()
    assert client.get(f"/api/projects/{case['project']}", headers=headers).status_code == expected
    assert client.put(f"/api/projects/{case['project']}", headers=headers,
                      json={'name': '不应写入', 'manager_id': case['users']['outsider']['id']}).status_code == expected
    response = client.get('/api/projects', headers=headers)
    if actor == 'guest':
        assert response.status_code == 403
    else:
        assert response.status_code == 200
        assert case['project'] not in {p['id'] for p in response.json()['items']}
        assert case['secret'] not in response.text
    assert security_snapshot() == before


@pytest.mark.parametrize('actor', ['outsider', 'finance_outsider', 'guest'])
def test_financial_reads_writes_and_replay_require_project_access(security_case, actor):
    case = security_case
    client, headers = case['client'], case['headers'][actor]
    expected = 403 if actor == 'guest' else 404
    before = security_snapshot()
    for url in [f"/api/receipts?project_id={case['project']}",
                f"/api/finance/receipts?project_id={case['project']}",
                f"/api/receipts/project/{case['project']}/stats",
                f"/api/contracts?project_id={case['project']}",
                f"/api/invoices?contract_id={case['contract']}",
                f"/api/payment-plans?contract_id={case['contract']}"]:
        response = client.get(url, headers=headers)
        assert response.status_code == expected, (url, response.status_code)
        assert case['secret'] not in response.text
    for method, url, body in [
        ('POST', '/api/receipts', receipt_body(case['project'])),
        ('POST', '/api/finance/receipts', case['receipt_body']),  # existing key: no early-return bypass
        ('POST', '/api/finance/receipts', {**case['receipt_body'], 'idempotency_key': uuid4().hex}),
        ('POST', '/api/contracts', {'project_id': case['project'], 'code': uuid4().hex, 'name': '不应创建', 'amount': '10.00'}),
        ('POST', '/api/payment-plans', {'contract_id': case['contract'], 'name': '不应创建', 'planned_amount': '10.00'}),
        ('POST', '/api/invoices', {'contract_id': case['contract'], 'code': uuid4().hex, 'amount': '10.00'}),
        ('PUT', f"/api/invoices/{case['invoice']}/red-invoice", None),
    ]:
        response = client.request(method, url, headers=headers, json=body)
        assert response.status_code == expected, (url, response.status_code)
    for url in ['/api/receipts', '/api/finance/receipts', '/api/contracts', '/api/invoices']:
        response = client.get(url, headers=headers)
        if actor == 'guest':
            assert response.status_code == 403
        else:
            assert response.status_code == 200
            assert case['receipt'] not in {row['id'] for row in response.json()}
            assert case['secret'] not in response.text
    assert security_snapshot() == before


@pytest.mark.parametrize('actor', ['member', 'outsider', 'finance_outsider', 'guest'])
def test_both_reversal_actions_and_confirmation_deny_without_effects(security_case, actor):
    case = security_case
    client, headers = case['client'], case['headers'][actor]
    expected = 404 if actor == 'finance_outsider' else 403
    before = security_snapshot()
    for method, url in [('PUT', f"/api/receipts/{case['receipt']}/confirm"),
                        ('PUT', f"/api/receipts/{case['receipt']}/reverse"),
                        ('POST', f"/api/finance/receipts/{case['receipt']}/write-off")]:
        response = client.request(method, url, headers=headers, json={'reason': '不应冲销'})
        assert response.status_code == expected, (url, response.status_code)
    assert security_snapshot() == before


@pytest.mark.parametrize('actor', ['owner', 'member', 'admin', 'enterprise_admin'])
def test_authorized_project_read_and_update_still_work(security_case, actor):
    case = security_case
    client, headers = case['client'], case['headers'][actor]
    response = client.get(f"/api/projects/{case['project']}", headers=headers)
    assert response.status_code == 200
    assert response.json()['name'] == case['secret']
    assert client.put(f"/api/projects/{case['project']}", headers=headers, json={'description': '合法更新'}).status_code == 200
    from app.models import Project
    with SessionLocal() as db:
        assert db.get(Project, case['project']).description == '合法更新'
    assert client.get(f"/api/projects/{case['project']}", headers=headers).json()['name'] == case['secret']
    assert client.get(f"/api/receipts/project/{case['project']}/stats", headers=headers).status_code == 200
    assert client.get(f"/api/contracts?project_id={case['project']}", headers=headers).json()[0]['id'] == case['contract']
    if actor == 'owner':
        from app.models import ProjectMember
        with SessionLocal() as db:
            assert db.query(ProjectMember).filter_by(project_id=case['project'], user_id=case['users']['owner']['id']).count() == 0


@pytest.mark.parametrize('actor', ['finance_member', 'owner', 'admin', 'enterprise_admin'])
def test_financial_role_and_project_scope_positive_preserves_money_and_audit(security_case, actor):
    case = security_case
    client, headers = case['client'], case['headers'][actor]
    assert client.put(f"/api/receipts/{case['receipt']}/confirm", headers=headers).status_code == 200
    assert client.get(f"/api/receipts/project/{case['project']}/stats", headers=headers).json()['total_received'] == 25
    assert client.post(f"/api/finance/receipts/{case['receipt']}/write-off", headers=headers,
                       json={'reason': '合法合成冲销'}).status_code == 200
    before = security_snapshot()
    assert client.post(f"/api/finance/receipts/{case['receipt']}/write-off", headers=headers,
                       json={'reason': '不应重复冲销'}).status_code == 400
    assert security_snapshot() == before
    assert client.get(f"/api/receipts/project/{case['project']}/stats", headers=headers).json()['total_received'] == 0
    with SessionLocal() as db:
        receipt = db.get(Receipt, case['receipt'])
        assert (receipt.amount, receipt.status, receipt.version) == (Decimal('25.00'), 'reversed', 2)
        actions = [row.action for row in db.query(AuditLog).filter_by(target_id=case['receipt']).all()]
        assert sorted(actions) == ['confirm', 'create', 'write_off']


def test_membership_revocation_is_checked_before_both_idempotent_returns(security_case):
    from app.models import ProjectMember
    case = security_case
    client, headers = case['client'], case['headers']['member']
    own = client.post('/api/projects', headers=case['headers']['admin'], json={
        'code': uuid4().hex, 'name': '合成成员自有项目', 'customer_id': client.get(f"/api/projects/{case['project']}", headers=headers).json()['customer_id'],
        'manager_id': case['users']['member']['id'],
    })
    assert own.status_code == 200
    body = receipt_body(case['project'], idempotency_key=uuid4().hex)
    first = client.post('/api/finance/receipts', headers=headers, json=body)
    assert first.status_code == 201
    with SessionLocal() as db:
        db.query(ProjectMember).filter_by(project_id=case['project'], user_id=case['users']['member']['id']).delete()
        db.commit()
    before = security_snapshot()
    assert client.post('/api/finance/receipts', headers=headers, json=body).status_code == 404
    # Authorized requested project cannot disguise a key pointing at a forbidden receipt.
    assert client.post('/api/finance/receipts', headers=headers,
                       json={**body, 'project_id': own.json()['id']}).status_code == 404
    assert security_snapshot() == before
    with SessionLocal() as db:
        db.add(ProjectMember(project_id=case['project'], user_id=case['users']['member']['id']))
        db.commit()
    replay = client.post('/api/finance/receipts', headers=headers, json=body)
    assert replay.status_code == 200
    assert replay.json()['id'] == first.json()['id']


def test_related_channels_cannot_disclose_or_modify_foreign_project(security_case):
    case = security_case
    client, headers = case['client'], case['headers']['outsider']
    before = security_snapshot()
    for url in [f"/api/tasks?project_id={case['project']}", f"/api/milestones?project_id={case['project']}",
                f"/api/gantt/{case['project']}", f"/api/calendar?project_id={case['project']}",
                f"/api/timesheets?project_id={case['project']}", f"/api/resource-load?project_id={case['project']}",
                f"/api/reports/budget-vs-actual?project_id={case['project']}",
                f"/api/export/receipts?project_id={case['project']}",
                f"/api/documents/{case['document']}/versions", f"/api/files/list?project_id={case['project']}"]:
        response = client.get(url, headers=headers)
        assert response.status_code == 404, (url, response.status_code)
    for method, url, body in [
        ('POST', '/api/tasks', {'project_id': case['project'], 'title': '不应创建'}),
        ('PUT', f"/api/tasks/{case['task']}", {'status': 'done'}),
        ('DELETE', f"/api/tasks/{case['task']}", None),
        ('PUT', '/api/tasks/batch-complete', {'ids': [case['task']]}),
        ('POST', '/api/milestones', {'project_id': case['project'], 'name': '不应创建'}),
        ('DELETE', f"/api/documents/{case['document']}", None),
        ('POST', '/api/timesheets', {'project_id': case['project'], 'hours': '1', 'work_date': '2026-01-15'}),
    ]:
        assert client.request(method, url, headers=headers, json=body).status_code == 404
    search = client.get('/api/search', headers=headers, params={'q': case['secret']})
    assert search.status_code == 200
    assert all(not search.json()[key] for key in ['projects', 'tasks', 'contracts', 'risks', 'issues', 'receipts', 'milestones'])
    for url in ['/api/reports/budget-vs-actual', '/api/reports/cost-profit', '/api/workbench', '/api/dashboard', '/api/audit-logs']:
        response = client.get(url, headers=headers)
        assert response.status_code == 200
        assert case['secret'] not in response.text, url
    from io import BytesIO
    import openpyxl
    for url in ['/api/export/projects', '/api/export/receipts', '/api/export/contracts', '/api/export/aging']:
        response = client.get(url, headers=headers)
        assert response.status_code == 200
        if 'spreadsheet' in response.headers.get('content-type', ''):
            workbook = openpyxl.load_workbook(BytesIO(response.content), read_only=True)
            values = [str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
            assert not any(case['secret'] in value for value in values)
        else:
            assert case['secret'] not in response.text
    assert security_snapshot() == before


def test_denied_upload_download_delete_has_no_file_or_database_effect(security_case):
    import hashlib
    from app.routes.files import UPLOAD_DIR
    case = security_case
    client, member = case['client'], case['headers']['member']
    # Create actual synthetic files so an unauthorized 404 cannot be a missing-file false positive.
    attachment = client.post(f"/api/files/upload?project_id={case['project']}", headers=member,
                             files={'file': ('synthetic.txt', b'synthetic attachment', 'text/plain')})
    assert attachment.status_code == 200
    version = client.post(f"/api/documents/{case['document']}/versions", headers=member,
                          files={'file': ('synthetic-version.txt', b'synthetic version', 'text/plain')})
    assert version.status_code == 200
    file_url = f"/api/files/download/{case['project']}/{attachment.json()['id']}"
    version_url = f"/api/documents/{case['document']}/versions/{version.json()['id']}/download"
    assert client.get(file_url, headers=member).content == b'synthetic attachment'
    assert client.get(version_url, headers=member).content == b'synthetic version'

    def files():
        return {str(path.relative_to(UPLOAD_DIR)): hashlib.sha256(path.read_bytes()).hexdigest()
                for path in UPLOAD_DIR.rglob('*') if path.is_file()}

    before_rows, before_files = security_snapshot(), files()
    for actor, expected in [('outsider', 404), ('guest', 403)]:
        headers = case['headers'][actor]
        assert client.get(file_url, headers=headers).status_code == expected
        assert client.get(version_url, headers=headers).status_code == expected
        assert client.delete(f"/api/files/{case['project']}/{attachment.json()['id']}", headers=headers).status_code == expected
        assert client.delete(f"/api/documents/{case['document']}", headers=headers).status_code == expected
        assert client.post(f"/api/files/upload?project_id={case['project']}", headers=headers,
                           files={'file': ('must-not-exist.txt', b'denied', 'text/plain')}).status_code == expected
        assert client.post(f"/api/documents/{case['document']}/versions", headers=headers,
                           files={'file': ('must-not-exist.txt', b'denied', 'text/plain')}).status_code == expected
    assert files() == before_files
    assert security_snapshot() == before_rows


def test_mixed_task_batch_is_rejected_before_any_task_changes(security_case):
    case = security_case
    client, headers = case['client'], case['headers']['outsider']
    own_task = client.post('/api/tasks', headers=headers,
                          json={'project_id': case['other_project'], 'title': '合成合法任务'})
    assert own_task.status_code == 200
    before = security_snapshot()
    response = client.put('/api/tasks/batch-complete', headers=headers,
                          json={'ids': [own_task.json()['id'], case['task']]})
    assert response.status_code == 404
    assert security_snapshot() == before


def test_member_cannot_reassign_scope_or_replace_project_relationships(security_case):
    case = security_case
    client, headers = case['client'], case['headers']['member']
    before = security_snapshot()
    assert client.put(f"/api/projects/{case['project']}", headers=headers,
                      json={'manager_id': case['users']['outsider']['id']}).status_code == 403
    assert client.put(f"/api/projects/{case['project']}", headers=headers,
                      json={'receipts': [], 'members': []}).status_code == 400
    assert security_snapshot() == before


def test_import_preflights_foreign_projects_before_any_rows_are_written(security_case):
    from io import BytesIO
    import openpyxl
    case = security_case
    client, headers = case['client'], case['headers']['outsider']
    target = client.get(f"/api/projects/{case['project']}", headers=case['headers']['admin']).json()
    own = client.get(f"/api/projects/{case['other_project']}", headers=headers).json()
    for kind, columns, rows in [
        ('projects', ['项目编号', '项目名称', '客户名称', '合同金额'],
         [[own['code'], '不应部分更新', '不应新增客户', '10'], [target['code'], '不应越权更新', '不应新增客户', '10']]),
        ('receipts', ['项目编号', '付款方', '回款金额', '回款日期'],
         [[own['code'], '不应部分入账', '10', '2026-01-15'], [target['code'], '不应越权入账', '10', '2026-01-15']]),
    ]:
        workbook = openpyxl.Workbook()
        workbook.active.append(columns)
        for row in rows:
            workbook.active.append(row)
        output = BytesIO()
        workbook.save(output)
        before = security_snapshot()
        response = client.post(f'/api/import/excel?import_type={kind}', headers=headers,
                               files={'file': ('synthetic.xlsx', output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        assert response.status_code == 404
        assert security_snapshot() == before


def test_project_summary_and_dashboard_cannot_include_foreign_money(security_case):
    case = security_case
    client, headers = case['client'], case['headers']['outsider']
    assert client.put(f"/api/receipts/{case['receipt']}/confirm", headers=case['headers']['admin']).status_code == 200
    before = security_snapshot()
    summary = client.get('/api/projects/summary', headers=headers)
    dashboard = client.get('/api/dashboard', headers=headers)
    assert summary.status_code == dashboard.status_code == 200
    assert summary.json()['total'] == 1
    assert summary.json()['completion_rate'] == 0
    assert summary.json()['budget_used_pct'] == 0
    assert dashboard.json()['total_projects'] == 1
    assert dashboard.json()['total_contract_amount'] == 0
    assert dashboard.json()['total_received'] == 0
    assert client.get('/api/dashboard/summary', headers=headers).json()['total_received'] == 0
    assert security_snapshot() == before


def test_idempotent_integrity_retry_rechecks_permission_after_rollback(security_case, monkeypatch):
    """Inject a stale key read, then use the real unique constraint/rollback path.

    SQLite serializes normal writers. This fault injection reaches the retry
    branch used after a constraint race and revokes membership after rollback;
    it does not mock the business result or replace the database.
    """
    from sqlalchemy.orm import Session
    from app.models import ProjectMember, ReceiptRequest
    case = security_case
    client, headers = case['client'], case['headers']['member']
    body = receipt_body(case['project'], amount='10.00', idempotency_key=uuid4().hex)
    original = client.post('/api/finance/receipts', headers=headers, json=body)
    assert original.status_code == 201
    before = security_snapshot()
    get_row, rollback = Session.get, Session.rollback
    state = {'hidden': False, 'revoked': False}

    def stale_get(db, entity, ident, *args, **kwargs):
        if entity is ReceiptRequest and ident == body['idempotency_key'] and not state['hidden']:
            state['hidden'] = True
            return None
        return get_row(db, entity, ident, *args, **kwargs)

    def revoke_after_rollback(db, *args, **kwargs):
        result = rollback(db, *args, **kwargs)
        if state['hidden'] and not state['revoked']:
            state['revoked'] = True
            with SessionLocal() as other:
                other.query(ProjectMember).filter_by(project_id=case['project'], user_id=case['users']['member']['id']).delete()
                other.commit()
        return result

    monkeypatch.setattr(Session, 'get', stale_get)
    monkeypatch.setattr(Session, 'rollback', revoke_after_rollback)
    response = client.post('/api/finance/receipts', headers=headers, json=body)
    assert response.status_code == 404
    assert state == {'hidden': True, 'revoked': True}
    after = security_snapshot()
    # The injected external membership change is expected; every business table
    # including receipt requests, allocations, audit and activity must be intact.
    assert {k: v for k, v in after.items() if k != 'project_members'} == {k: v for k, v in before.items() if k != 'project_members'}
    with SessionLocal() as db:
        assert db.get(Receipt, original.json()['id']).amount == Decimal('10.00')
        assert db.query(ReceiptRequest).filter_by(key=body['idempotency_key']).count() == 1


def test_project_import_cannot_bypass_finance_confirmation_permission(security_case):
    from io import BytesIO
    import openpyxl
    case = security_case
    client, headers = case['client'], case['headers']['outsider']
    own = client.get(f"/api/projects/{case['other_project']}", headers=headers).json()
    workbook = openpyxl.Workbook()
    workbook.active.append(['项目编号', '项目名称', '客户名称', '合同金额', '初始已收金额'])
    workbook.active.append([own['code'], '不应更新并确认回款', '不应新增客户', '100', '25'])
    output = BytesIO()
    workbook.save(output)
    before = security_snapshot()
    response = client.post('/api/import/excel?import_type=projects', headers=headers,
                           files={'file': ('synthetic.xlsx', output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 403
    assert security_snapshot() == before


def test_guest_cannot_create_project_and_member_can_only_create_owned_project(security_case):
    case = security_case
    client = case['client']
    customer_id = client.get(f"/api/projects/{case['project']}", headers=case['headers']['admin']).json()['customer_id']
    body = {'code': uuid4().hex, 'name': '合成新项目', 'customer_id': customer_id, 'manager_id': ''}
    before = security_snapshot()
    assert client.post('/api/projects', headers=case['headers']['guest'], json=body).status_code == 403
    assert client.post('/api/projects', headers=case['headers']['member'],
                       json={**body, 'manager_id': case['users']['outsider']['id']}).status_code == 403
    assert security_snapshot() == before
    created = client.post('/api/projects', headers=case['headers']['member'], json=body)
    assert created.status_code == 200
    assert created.json()['manager_id'] == case['users']['member']['id']
    assert client.get(f"/api/projects/{created.json()['id']}", headers=case['headers']['member']).status_code == 200
    assert client.get(f"/api/projects/{created.json()['id']}", headers=case['headers']['outsider']).status_code == 404


@pytest.mark.parametrize('action', ['confirm', 'reverse', 'write-off'])
def test_finance_action_role_is_rechecked_at_write_transaction(security_case, monkeypatch, action):
    from app.routes import finance
    case = security_case
    client, headers = case['client'], case['headers']['finance_member']
    if action != 'confirm':
        assert client.put(f"/api/receipts/{case['receipt']}/confirm", headers=case['headers']['admin']).status_code == 200
    before = security_snapshot()
    original_lock = finance._lock_receipt_write

    def downgrade_before_write_lock(db):
        # Simulate a concurrent administrator revoking the action role after
        # the request dependency was evaluated, before the write transaction.
        with SessionLocal() as other:
            other.get(User, case['users']['finance_member']['id']).role = 'member'
            other.commit()
        original_lock(db)

    monkeypatch.setattr(finance, '_lock_receipt_write', downgrade_before_write_lock)
    if action == 'write-off':
        response = client.post(f"/api/finance/receipts/{case['receipt']}/write-off", headers=headers, json={'reason': '不应冲销'})
    else:
        response = client.put(f"/api/receipts/{case['receipt']}/{action}", headers=headers)
    assert response.status_code == 403
    after = security_snapshot()
    assert {k: v for k, v in after.items() if k != 'users'} == {k: v for k, v in before.items() if k != 'users'}


def test_concurrent_confirmation_adds_one_audit(archive_case):
    _, headers, project_id, _, _ = archive_case
    with TestClient(app, client=(f"create-{uuid4().hex}", 50000)) as client:
        receipt = client.post("/api/receipts", headers=headers, json=receipt_body(project_id))
    assert receipt.status_code == 200
    receipt_id = receipt.json()["id"]
    barrier = Barrier(2)

    def confirm(_):
        with TestClient(app, client=(f"confirm-{uuid4().hex}", 50000)) as client:
            barrier.wait(timeout=10)
            return client.put(f"/api/receipts/{receipt_id}/confirm", headers=headers)

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(confirm, range(2)))
    assert sorted(r.status_code for r in responses) == [200, 400]
    with SessionLocal() as db:
        assert db.query(AuditLog).filter_by(target_type="receipt", target_id=receipt_id, action="confirm").count() == 1


def test_demo_schema_addition_preserves_existing_financial_rows(tmp_path):
    """Model delta is one new table; create_all must not rebuild old tables."""
    from datetime import datetime
    from sqlalchemy import create_engine, inspect, text
    from sqlalchemy.orm import Session
    from app.database import Base
    from app.models import Customer, Project, ReceiptRequest

    engine = create_engine(f"sqlite:///{tmp_path / 'synthetic-old-schema.db'}")
    try:
        old_tables = [table for table in Base.metadata.tables.values() if table is not ReceiptRequest.__table__]
        Base.metadata.create_all(bind=engine, tables=old_tables)
        before_tables = set(inspect(engine).get_table_names())
        with Session(engine) as db:
            db.add(User(id='upgrade-user', username='synthetic-upgrade', hashed_password='synthetic-only', real_name='合成用户'))
            db.add(Customer(id='upgrade-customer', name='虚构旧库客户'))
            db.flush()
            db.add(Project(id='upgrade-project', code='SYNTHETIC-OLD', name='虚构旧库项目',
                           customer_id='upgrade-customer', manager_id='upgrade-user'))
            db.flush()
            db.add(Receipt(id='upgrade-receipt', project_id='upgrade-project', amount=Decimal('25.00'),
                           received_date=datetime(2026, 1, 15), status='confirmed', notes='保留的合成历史备注'))
            db.add(AuditLog(id='upgrade-audit', user_id='upgrade-user', action='create', target_type='receipt',
                            target_id='upgrade-receipt', new_value='25.00'))
            db.commit()
        with engine.connect() as connection:
            before_receipts = connection.execute(text('SELECT * FROM receipts')).all()
            before_audit = connection.execute(text('SELECT * FROM audit_logs')).all()
        Base.metadata.create_all(bind=engine)
        assert set(inspect(engine).get_table_names()) - before_tables == {'receipt_requests'}
        with engine.connect() as connection:
            assert connection.execute(text('SELECT * FROM receipts')).all() == before_receipts
            assert connection.execute(text('SELECT * FROM audit_logs')).all() == before_audit
            assert connection.execute(text('PRAGMA foreign_key_check')).all() == []
            assert connection.execute(text('PRAGMA integrity_check')).scalar() == 'ok'
    finally:
        engine.dispose()


def test_legacy_negative_allocation_cannot_inflate_available_balance(archive_case):
    from datetime import datetime
    client, headers, project_id, invoice_id, _ = archive_case
    # Simulate a row accepted before strict validation, without rewriting history.
    with SessionLocal() as db:
        receipt = Receipt(project_id=project_id, amount=Decimal('-10.00'),
                          received_date=datetime(2026, 1, 15), status='confirmed')
        db.add(receipt)
        db.flush()
        db.add(ReceiptAllocation(receipt_id=receipt.id, invoice_id=invoice_id, allocated_amount=Decimal('-10.00')))
        db.commit()
    response = client.post('/api/finance/receipts', headers=headers, json=receipt_body(
        project_id, amount='100.00', idempotency_key=uuid4().hex,
        allocations=[{'invoice_id': invoice_id, 'allocated_amount': '100.00'}],
    ))
    assert response.status_code == 400
    assert receipt_count(project_id) == 1
    with SessionLocal() as db:
        assert db.query(ReceiptAllocation).filter_by(invoice_id=invoice_id).one().allocated_amount == Decimal('-10.00')


def test_normalized_amount_and_payment_date_alias_replay_same_request(archive_case):
    client, headers, project_id, _, _ = archive_case
    key = uuid4().hex
    body = receipt_body(project_id, amount='25.0', idempotency_key=key)
    first = client.post('/api/finance/receipts', headers=headers, json=body)
    assert first.status_code == 201
    equivalent = {**body, 'amount': '25.00', 'payment_date': body['received_date']}
    del equivalent['received_date']
    replay = client.post('/api/finance/receipts', headers=headers, json=equivalent)
    assert replay.status_code == 200
    assert replay.json()['id'] == first.json()['id']
    conflict = client.post('/api/finance/receipts', headers=headers,
                           json={**body, 'payment_date': '2026-01-16T00:00:00'})
    assert conflict.status_code == 422
    assert receipt_count(project_id) == 1
